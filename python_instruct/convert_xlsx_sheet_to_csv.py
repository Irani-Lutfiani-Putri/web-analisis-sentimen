"""Convert one sheet of an .xlsx workbook to a UTF-8 CSV file.

Usage:
    python python_instruct/convert_xlsx_sheet_to_csv.py <workbook.xlsx> <SheetName> <output.csv> [--list]

Writes UTF-8 with BOM (utf-8-sig) because the app decodes uploads with
utf-8-sig and Excel opens BOM-marked CSV correctly. Embedded newlines and
quotes are handled by the csv module (RFC 4180 quoting), so multi-line
comments survive the round trip.

Run with --list to print sheet names and row counts without converting.
"""
import csv
import sys
from pathlib import Path


def load(path):
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl belum terpasang. Jalankan dengan:\n"
                 "  uv run --no-project --python 3.11 --with openpyxl python python_instruct/convert_xlsx_sheet_to_csv.py ...")
    return openpyxl.load_workbook(path, read_only=True, data_only=True)


def main(argv):
    if len(argv) < 2:
        sys.exit(__doc__)
    workbook_path = Path(argv[1])
    if not workbook_path.is_file():
        sys.exit(f"Berkas tidak ditemukan: {workbook_path}")

    workbook = load(workbook_path)
    try:
        if "--list" in argv:
            for index, name in enumerate(workbook.sheetnames, 1):
                sheet = workbook[name]
                print(f"{index}. {name}  rows={sheet.max_row} cols={sheet.max_column}")
            return

        if len(argv) < 4:
            sys.exit(__doc__)
        sheet_name, output_path = argv[2], Path(argv[3])
        if sheet_name not in workbook.sheetnames:
            sys.exit(f"Sheet '{sheet_name}' tidak ada. Tersedia: {', '.join(workbook.sheetnames)}")

        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            sys.exit("Sheet kosong")

        output_path.parent.mkdir(parents=True, exist_ok=True)
        written = 0
        skipped = 0
        with output_path.open("w", encoding="utf-8-sig", newline="") as stream:
            writer = csv.writer(stream, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")
            writer.writerow(["" if cell is None else str(cell) for cell in header])
            for row in rows:
                if row is None or all(cell is None or str(cell).strip() == "" for cell in row):
                    skipped += 1
                    continue
                writer.writerow(["" if cell is None else str(cell) for cell in row])
                written += 1

        print(f"sheet      : {sheet_name}")
        print(f"kolom      : {[str(c) for c in header]}")
        print(f"baris tulis: {written}")
        print(f"baris kosong dilewati: {skipped}")
        print(f"output     : {output_path}  ({output_path.stat().st_size:,} byte)")
    finally:
        workbook.close()


if __name__ == "__main__":
    main(sys.argv)
